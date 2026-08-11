from __future__ import annotations

from pathlib import Path

from pptx.enum.text import PP_ALIGN

from .components import (
    add_big_number,
    add_bullets,
    add_card,
    add_footer,
    add_image,
    add_light_card,
    add_note,
    add_section_label,
    add_table,
    add_text,
    add_title,
    blank_slide,
    prs_new,
    set_bg,
)
from .config import TEMP_PPTX_PATH
from .metrics import fmt_num, fmt_pct, fmt_rank
from .storytelling import attention_points, driver_flow, executive_highlights, phrase_ranking, phrase_variation, priorities, title_for_stage


def row_stage(metrics: dict, stage: str):
    return metrics["rosario"][metrics["rosario"]["etapa"] == stage].iloc[0]


def add_standard_slide(prs, title: str, subtitle: str | None = None):
    slide = blank_slide(prs)
    set_bg(slide)
    add_title(slide, title, subtitle)
    add_footer(slide)
    return slide


def add_title_slide(prs):
    slide = blank_slide(prs)
    set_bg(slide)
    add_text(slide, "IDEB 2025", 0.82, 1.28, 5.5, 0.65, size=36, color="navy", bold=True)
    add_text(slide, "Rosario do Catete", 0.82, 1.95, 6.2, 0.68, size=34, color="navy", bold=True)
    add_text(slide, "Resultados, evolucao e posicionamento da Rede Municipal", 0.86, 2.72, 6.3, 0.35, size=16, color="dark_text")
    add_text(slide, "Comparativo 2023 x 2025", 0.86, 3.18, 3.5, 0.3, size=13, color="cyan", bold=True)
    add_card(slide, "Rede", "Municipal", "Ensino Fundamental", 0.88, 4.35, 2.15, 1.18, "cyan")
    add_card(slide, "Periodo", "2023-2025", "Evolucao do IDEB", 3.18, 4.35, 2.15, 1.18, "navy")
    add_card(slide, "Foco", "DRE 4", "Ranking regional", 5.48, 4.35, 2.15, 1.18, "magenta")
    add_light_card(slide, "Entrega executiva", "Apresentacao automatizada com dados oficiais, rankings, graficos e diagnostico gerencial.", 8.35, 1.3, 3.8, 4.2, "yellow")
    add_footer(slide)


def add_overview_slide(prs):
    slide = add_standard_slide(prs, "O que vamos analisar", "Cinco dimensoes para transformar dados em decisao")
    cards = [
        ("1", "Evolucao", "IDEB 2023 x 2025"),
        ("2", "Comparacao", "Sergipe municipal e DRE 4"),
        ("3", "Ranking", "Posicao estadual e regional"),
        ("4", "Aprendizagem e fluxo", "Origem da variacao"),
        ("5", "Escolas", "Quem avancou e onde agir"),
    ]
    colors = ["cyan", "navy", "magenta", "yellow", "aqua"]
    for i, (num, title, body) in enumerate(cards):
        x = 0.75 + i * 2.45
        add_card(slide, title, num, body, x, 2.3, 2.1, 1.45, colors[i], "dark_text" if colors[i] == "yellow" else "white")
    add_note(slide, "A narrativa segue a sequencia: onde estavamos, onde estamos, como nos comparamos, por que mudou e onde agir.", 1.55, 5.0, 10.2, 0.72)


def add_summary_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Resumo executivo", "Rosario avancou nas duas etapas avaliadas")
    ini = row_stage(metrics, "Anos iniciais")
    fin = row_stage(metrics, "Anos finais")
    add_card(slide, "IDEB 2025", fmt_num(ini.ideb_2025), "Anos iniciais", 0.75, 1.72, 2.25, 1.25, "cyan")
    add_card(slide, "IDEB 2025", fmt_num(fin.ideb_2025), "Anos finais", 3.2, 1.72, 2.25, 1.25, "navy")
    add_card(slide, "Variacao", fmt_num(ini.variacao_absoluta, signed=True), "Maior avanco", 5.65, 1.72, 2.25, 1.25, "aqua")
    add_card(slide, "Ranking estadual", fmt_rank(ini.ranking_estadual_seq_2025), "Melhor posicao 2025", 8.1, 1.72, 2.25, 1.25, "magenta")
    add_card(slide, "Ranking DRE 4", fmt_rank(ini.ranking_dre4_seq_2025), "Melhor posicao 2025", 10.55, 1.72, 2.25, 1.25, "yellow", "dark_text")
    add_image(slide, metrics["charts"]["resumo"], 1.0, 3.35, 7.3, 2.95)
    add_note(slide, "Leitura executiva: a rede municipal cresceu em IDEB nas duas etapas, com ganho regional mais forte nos anos iniciais.", 8.55, 3.58, 3.55, 1.5)


def add_what_happened_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Rosario avancou entre 2023 e 2025", "O movimento foi positivo nos anos iniciais e finais")
    ini = row_stage(metrics, "Anos iniciais")
    fin = row_stage(metrics, "Anos finais")
    add_big_number(slide, fmt_num(ini.ideb_2023), "Anos iniciais em 2023", 1.05, 2.05, "navy")
    add_text(slide, "->", 3.05, 2.1, 0.8, 0.5, size=30, color="mid_gray", bold=True, align=PP_ALIGN.CENTER)
    add_big_number(slide, fmt_num(ini.ideb_2025), "Anos iniciais em 2025", 4.0, 2.05, "cyan")
    add_big_number(slide, fmt_num(fin.ideb_2023), "Anos finais em 2023", 1.05, 4.0, "navy")
    add_text(slide, "->", 3.05, 4.05, 0.8, 0.5, size=30, color="mid_gray", bold=True, align=PP_ALIGN.CENTER)
    add_big_number(slide, fmt_num(fin.ideb_2025), "Anos finais em 2025", 4.0, 4.0, "cyan")
    add_note(slide, phrase_variation(ini) + "\n" + phrase_variation(fin), 7.45, 2.25, 4.45, 2.4)


def add_stage_slide(prs, metrics: dict, stage: str, chart_key: str):
    row = row_stage(metrics, stage)
    slide = add_standard_slide(prs, title_for_stage(row), f"{stage}: evolucao 2023 x 2025")
    add_image(slide, metrics["charts"][chart_key], 0.72, 1.72, 8.55, 3.15)
    add_card(slide, "Variacao", fmt_num(row.variacao_absoluta, signed=True), "pontos de IDEB", 9.55, 1.78, 2.55, 1.18, "aqua" if row.variacao_absoluta >= 0 else "magenta")
    add_card(slide, "Crescimento", fmt_pct(row.variacao_percentual), "2023-2025", 9.55, 3.17, 2.55, 1.18, "cyan")
    add_note(slide, phrase_ranking(row, "dre4"), 9.25, 4.78, 3.05, 0.92)


def add_stage_comparison_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Anos iniciais lideram o crescimento", "Comparacao direta entre etapas")
    add_image(slide, metrics["charts"]["resumo"], 0.85, 1.65, 7.45, 3.7)
    rows = [["Etapa", "IDEB 2023", "IDEB 2025", "Var."]]
    for _, row in metrics["rosario"].sort_values("etapa").iterrows():
        rows.append([row.etapa, fmt_num(row.ideb_2023), fmt_num(row.ideb_2025), fmt_num(row.variacao_absoluta, signed=True)])
    add_table(slide, rows, 8.55, 2.0, 3.45, 1.55)
    add_note(slide, "A escala dos graficos de IDEB preserva o intervalo 0 a 10 para evitar distorcoes visuais.", 8.55, 4.15, 3.45, 0.85)


def add_sergipe_slide(prs, metrics: dict, stage: str, chart_key: str):
    slide = add_standard_slide(prs, sergipe_title(metrics, stage), "Comparacao com a media municipal dos municipios sergipanos disponiveis")
    add_image(slide, metrics["charts"][chart_key], 0.95, 1.65, 7.35, 4.25)
    row = row_stage(metrics, stage)
    media = metrics["media_sergipe"][(metrics["media_sergipe"]["etapa"] == stage) & (metrics["media_sergipe"]["ano"] == 2025)]["ideb"].iloc[0]
    diff = row.ideb_2025 - media
    msg = f"Em 2025, Rosario ficou {fmt_num(abs(diff))} ponto {'acima' if diff >= 0 else 'abaixo'} da media municipal de Sergipe."
    add_note(slide, msg + "\n\nNota: o agregado oficial UF 2025 foi retirado do escopo; esta leitura usa a media dos municipios.", 8.65, 2.15, 3.35, 2.0)


def add_dre4_position_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, dre4_position_title(metrics), "Ranking regional de 2025")
    add_image(slide, metrics["charts"]["dre4_iniciais"], 0.7, 1.55, 5.85, 4.65)
    add_image(slide, metrics["charts"]["dre4_finais"], 6.82, 1.55, 5.85, 4.65)
    add_section_label(slide, "Iniciais", 1.0, 1.32, "cyan")
    add_section_label(slide, "Finais", 7.1, 1.32, "magenta")


def add_dre4_evolution_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "A DRE 4 tambem mudou de patamar", "Variacao de IDEB entre 2023 e 2025")
    add_image(slide, metrics["charts"]["dre4_var_iniciais"], 0.7, 1.55, 5.85, 4.6)
    add_image(slide, metrics["charts"]["dre4_var_finais"], 6.82, 1.55, 5.85, 4.6)


def add_dre4_top_variation_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Rosario aparece entre os avancos da DRE 4", "Variacao absoluta do IDEB entre 2023 e 2025")
    add_image(slide, metrics["charts"]["dre4_var_iniciais"], 0.7, 1.55, 5.85, 4.65)
    add_image(slide, metrics["charts"]["dre4_var_finais"], 6.82, 1.55, 5.85, 4.65)
    add_section_label(slide, "Iniciais", 1.0, 1.32, "cyan")
    add_section_label(slide, "Finais", 7.1, 1.32, "magenta")


def add_state_ranking_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, state_ranking_title(metrics), "Ranking sequencial entre municipios com IDEB disponivel")
    add_image(slide, metrics["charts"]["ranking_estadual"], 0.95, 1.72, 6.7, 4.35)
    ini = row_stage(metrics, "Anos iniciais")
    fin = row_stage(metrics, "Anos finais")
    add_note(slide, phrase_ranking(ini, "estadual"), 8.05, 2.05, 3.75, 1.1)
    add_note(slide, phrase_ranking(fin, "estadual"), 8.05, 3.6, 3.75, 1.1)
    add_text(slide, "Ranking sequencial: empates sao ordenados alfabeticamente, sem alterar notas.", 8.08, 5.35, 3.75, 0.35, size=8, color="mid_gray")


def add_state_context_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Recorte do ranking estadual", "Municipios proximos a Rosario em 2025")
    add_image(slide, metrics["charts"]["contexto_estadual_iniciais"], 0.7, 1.55, 5.85, 4.65)
    add_image(slide, metrics["charts"]["contexto_estadual_finais"], 6.82, 1.55, 5.85, 4.65)


def add_flow_concept_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Aprendizagem e fluxo explicam o IDEB", "IDEB combina desempenho no SAEB e rendimento escolar")
    ini = row_stage(metrics, "Anos iniciais")
    fin = row_stage(metrics, "Anos finais")
    add_card(slide, "Desempenho", fmt_num(ini.nota_saeb_media_padronizada_2025), "Anos iniciais 2025", 1.0, 2.05, 2.55, 1.35, "navy")
    add_card(slide, "Fluxo", fmt_num(ini.indicador_rendimento_2025, 3), "Anos iniciais 2025", 4.0, 2.05, 2.55, 1.35, "cyan")
    add_card(slide, "Desempenho", fmt_num(fin.nota_saeb_media_padronizada_2025), "Anos finais 2025", 7.0, 2.05, 2.55, 1.35, "magenta")
    add_card(slide, "Fluxo", fmt_num(fin.indicador_rendimento_2025, 3), "Anos finais 2025", 10.0, 2.05, 2.55, 1.35, "yellow", "dark_text")
    add_note(slide, "Leitura: quando desempenho e fluxo crescem juntos, o IDEB tende a avancar com mais consistencia.", 2.2, 4.65, 8.9, 0.85)


def add_flow_diagnostic_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "O avanco veio de desempenho e fluxo", "Diagnostico da origem da variacao")
    add_image(slide, metrics["charts"]["fluxo_desempenho"], 0.9, 1.6, 6.8, 4.55)
    msgs = [driver_flow(row_stage(metrics, "Anos iniciais")), driver_flow(row_stage(metrics, "Anos finais"))]
    add_note(slide, "\n".join(msgs), 8.1, 2.2, 3.75, 1.6)


def add_schools_overview_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Todas as escolas comparaveis avancaram", "Visao municipal por escola e etapa")
    count = metrics["contagem_escolas"]
    add_card(slide, "Resultados", str(count["comparaveis"]), "comparaveis", 0.9, 1.75, 2.2, 1.25, "navy")
    add_card(slide, "Avancaram", str(count["avancaram"]), "resultados escolares", 3.35, 1.75, 2.2, 1.25, "aqua")
    add_card(slide, "Estaveis", str(count["estaveis"]), "resultados escolares", 5.8, 1.75, 2.2, 1.25, "yellow", "dark_text")
    add_card(slide, "Recuaram", str(count["recuaram"]), "resultados escolares", 8.25, 1.75, 2.2, 1.25, "magenta")
    rows = [["Escola", "Etapa", "2023", "2025", "Var."]]
    for _, row in metrics["escolas_validas"].sort_values("variacao_absoluta", ascending=False).iterrows():
        rows.append([short_school(row.escola), row.etapa.replace("Anos ", ""), fmt_num(row.ideb_2023), fmt_num(row.ideb_2025), fmt_num(row.variacao_absoluta, signed=True)])
    add_table(slide, rows, 0.95, 3.55, 11.45, 2.25)


def add_school_ideb_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "IDEB por escola mostra avancos disseminados", "Dumbbell chart: 2023 x 2025")
    add_image(slide, metrics["charts"]["escolas_iniciais"], 0.65, 1.55, 5.95, 4.8)
    add_image(slide, metrics["charts"]["escolas_finais"], 6.88, 1.55, 5.95, 4.8)


def add_school_top_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Escolas que mais avancaram", "Ranking pela variacao absoluta do IDEB")
    add_image(slide, metrics["charts"]["variacao_escolas"], 0.8, 1.45, 7.4, 5.0)
    best = metrics["maior_avanco_escola"].iloc[0]
    add_card(slide, "Maior avanco", fmt_num(best.variacao_absoluta, signed=True), best.etapa, 8.65, 2.0, 2.8, 1.35, "aqua")
    add_note(slide, short_school(best.escola, 90), 8.65, 3.75, 3.15, 0.9)


def add_attention_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Pontos de atencao", "Leitura gerencial sem linguagem acusatoria")
    add_bullets(slide, attention_points(metrics), 1.0, 1.85, 5.6, 3.3, size=14)
    add_note(slide, "Mesmo com avanco geral, a gestao deve observar rankings, diferencas entre etapas e sustentacao dos resultados escolares.", 7.05, 2.15, 4.8, 1.7)


def add_positive_highlights_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Destaques positivos", "Evidencias que sustentam a narrativa de avanco")
    highlights = executive_highlights(metrics)
    colors = ["cyan", "navy", "magenta", "yellow", "aqua"]
    for i, item in enumerate(highlights):
        x = 0.85 + (i % 2) * 5.9
        y = 1.7 + (i // 2) * 1.35
        add_light_card(slide, f"Destaque {i + 1}", item, x, y, 5.25, 1.28, colors[i % len(colors)])


def add_diagnosis_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Diagnostico executivo", "Avancos, atencoes e oportunidades")
    cols = [
        ("AVANCOS", executive_highlights(metrics)[:3], "aqua"),
        ("ATENCOES", attention_points(metrics)[:3], "magenta"),
        ("OPORTUNIDADES", priorities(metrics)[:3], "yellow"),
    ]
    for i, (title, items, color) in enumerate(cols):
        x = 0.75 + i * 4.1
        add_light_card(slide, title, "\n".join([f"- {item}" for item in items]), x, 1.72, 3.55, 4.15, color)


def add_priorities_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Prioridades para o proximo ciclo", "Recomendacoes conectadas aos resultados")
    add_bullets(slide, priorities(metrics), 1.0, 1.7, 7.2, 4.6, size=15)
    add_card(slide, "Foco", "2027", "proximo ciclo de monitoramento", 9.05, 2.25, 2.45, 1.3, "navy")
    add_note(slide, "A agenda recomendada combina aprendizagem, fluxo, apoio escolar e uso sistematico de evidencias.", 8.55, 4.1, 3.45, 1.2)


def add_conclusion_slide(prs, metrics: dict):
    slide = add_standard_slide(prs, "Rosario avancou e deve sustentar a trajetoria", "Conclusao executiva")
    ini = row_stage(metrics, "Anos iniciais")
    fin = row_stage(metrics, "Anos finais")
    add_card(slide, "Anos iniciais", fmt_num(ini.ideb_2025), f"{fmt_num(ini.variacao_absoluta, signed=True)} ponto", 1.0, 2.0, 2.65, 1.45, "cyan")
    add_card(slide, "Anos finais", fmt_num(fin.ideb_2025), f"{fmt_num(fin.variacao_absoluta, signed=True)} ponto", 4.2, 2.0, 2.65, 1.45, "navy")
    add_card(slide, "DRE 4", fmt_rank(ini.ranking_dre4_seq_2025), "melhor posicao regional", 7.4, 2.0, 2.65, 1.45, "magenta")
    add_note(slide, "Os resultados mostram avanco educacional, com destaque regional nos anos iniciais e necessidade de manter acompanhamento tecnico nos anos finais.", 1.25, 4.45, 10.5, 1.0)


def add_closing_slide(prs):
    slide = blank_slide(prs)
    set_bg(slide)
    add_text(slide, "Educacao baseada em evidencias", 1.0, 2.15, 7.9, 0.68, size=34, color="navy", bold=True)
    add_text(slide, "Transformar dados em decisoes para continuar avancando.", 1.05, 2.98, 7.2, 0.35, size=16, color="dark_text")
    add_card(slide, "Dados", "->", "Analise", 1.05, 4.3, 2.0, 1.05, "cyan")
    add_card(slide, "Analise", "->", "Decisao", 3.35, 4.3, 2.0, 1.05, "navy")
    add_card(slide, "Decisao", "->", "Aprendizagem", 5.65, 4.3, 2.0, 1.05, "aqua")
    add_footer(slide)


def short_school(name: str, max_len: int = 48) -> str:
    cleaned = str(name).replace("ESCOLA-MUNICIPAL-", "EM ").replace("ESCOLA MUNICIPAL ", "EM ").replace("ESCOLA-MUL-", "EM ")
    cleaned = cleaned.replace("-", " ").replace("  ", " ")
    return cleaned if len(cleaned) <= max_len else cleaned[: max_len - 1] + "..."



def sergipe_title(metrics: dict, stage: str) -> str:
    row = row_stage(metrics, stage)
    media = metrics["media_sergipe"][(metrics["media_sergipe"]["etapa"] == stage) & (metrics["media_sergipe"]["ano"] == 2025)]["ideb"].iloc[0]
    diff = row.ideb_2025 - media
    if diff >= 0.05:
        return f"{stage}: Rosario fica acima da media municipal de SE"
    if diff <= -0.05:
        return f"{stage}: Rosario fica abaixo da media municipal de SE"
    return f"{stage}: Rosario fica em linha com a media municipal de SE"


def dre4_position_title(metrics: dict) -> str:
    ini = row_stage(metrics, "Anos iniciais")
    fin = row_stage(metrics, "Anos finais")
    return f"Rosario ocupa {fmt_rank(ini.ranking_dre4_seq_2025)} nos iniciais e {fmt_rank(fin.ranking_dre4_seq_2025)} nos finais na DRE 4"


def state_ranking_title(metrics: dict) -> str:
    ini = row_stage(metrics, "Anos iniciais")
    fin = row_stage(metrics, "Anos finais")
    if ini.mudanca_ranking_estadual_seq > 0 and fin.mudanca_ranking_estadual_seq < 0:
        return "Ranking estadual melhora nos iniciais e exige atencao nos finais"
    if ini.mudanca_ranking_estadual_seq > 0 and fin.mudanca_ranking_estadual_seq > 0:
        return "Rosario ganha posicoes no ranking estadual"
    return "Ranking estadual revela trajetorias diferentes por etapa"
def create_presentation(metrics: dict, charts: dict[Path]) -> tuple[object, int]:
    metrics["charts"] = charts
    prs = prs_new()
    add_title_slide(prs)
    add_overview_slide(prs)
    add_summary_slide(prs, metrics)
    add_what_happened_slide(prs, metrics)
    add_stage_slide(prs, metrics, "Anos iniciais", "iniciais")
    add_stage_slide(prs, metrics, "Anos finais", "finais")
    add_stage_comparison_slide(prs, metrics)
    add_sergipe_slide(prs, metrics, "Anos iniciais", "sergipe_iniciais")
    add_sergipe_slide(prs, metrics, "Anos finais", "sergipe_finais")
    add_dre4_position_slide(prs, metrics)
    add_dre4_evolution_slide(prs, metrics)
    add_dre4_top_variation_slide(prs, metrics)
    add_state_ranking_slide(prs, metrics)
    add_state_context_slide(prs, metrics)
    add_flow_concept_slide(prs, metrics)
    add_flow_diagnostic_slide(prs, metrics)
    add_schools_overview_slide(prs, metrics)
    add_school_ideb_slide(prs, metrics)
    add_school_top_slide(prs, metrics)
    add_attention_slide(prs, metrics)
    add_positive_highlights_slide(prs, metrics)
    add_diagnosis_slide(prs, metrics)
    add_priorities_slide(prs, metrics)
    add_conclusion_slide(prs, metrics)
    add_closing_slide(prs)
    TEMP_PPTX_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(TEMP_PPTX_PATH)
    return prs, len(prs.slides)
